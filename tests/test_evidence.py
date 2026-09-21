import json
import unittest
from datetime import date, timedelta
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch

import pandas as pd
from streamlit.testing.v1 import AppTest

from evidence import (ACTIONS, Change, Customer, Event, Project, Source, add_assessment,
    contradictions, current_assessments, customer_frame, edit_source, freshness,
    get, import_project, load_demo, parse_customers, review_assessment,
    revise_hypothesis, save_customers, triangulation)
from evidence import hypothesis_thesis
from analytics import activity_report
from evidence import ActivityEvent
from analysis_tools import Proposals, accept_proposals, ai_proposals, extract_document



class WorkspaceAITests(unittest.TestCase):
    def client(self, bad=False):
        from ai_workspace import BatchResult, Findings
        client=Mock()
        def parse(**kw):
            payload=json.loads(kw["input"][1]["content"])
            if kw["text_format"] is BatchResult:
                items=[]
                for source in payload["sources"]:
                    label,text=next(iter(source["chunks"].items()))
                    items.append(dict(source_id=source["id"],evidence_kind=source["kind"],hypothesis_id="H1",chunk=label,
                        quote="invented quotation" if bad else text,stance="Unclear",rationale="Review the specific customer context.",limitations="Not conclusive."))
                result=BatchResult(assessments=items,note="Test draft")
            else:
                a=payload["links"][0]
                result=Findings(findings=[dict(area="Dashboard",target="All",conclusion="Evidence requires review.",assessment_ids=[a["id"]],limitations="Small sample.",next_test="Interview another customer.")])
            return SimpleNamespace(status="completed",output_parsed=result)
        client.responses.parse.side_effect=parse
        return client

    def test_batch_incremental_and_atomic(self):
        from ai_workspace import analyze_workspace, fingerprint
        p=load_demo(); before=p.model_dump()
        client=self.client()
        updated=analyze_workspace(p,"test","model",client=client)
        self.assertEqual(p.model_dump(),before)
        self.assertEqual(updated.decisions,p.decisions)
        self.assertEqual(updated.ai_report["sources_analyzed"],240)
        self.assertEqual(updated.ai_report["input_hash"],fingerprint(updated))
        new=[a for a in updated.assessments if a.id not in {a.id for a in p.assessments}]
        self.assertTrue(new)
        self.assertTrue(all(a.status=="Pending" for a in new))
        self.assertEqual(import_project(updated.model_dump_json().encode()).ai_report,updated.ai_report)
        count=client.responses.parse.call_count
        again=analyze_workspace(updated,"test","model",client=client)
        self.assertEqual(client.responses.parse.call_count-count,1)
        self.assertEqual(again.ai_report["batches"],0)
        again.customers[0].note="Changed interview note."
        self.assertNotEqual(again.ai_report["input_hash"],fingerprint(again))
        refreshed=analyze_workspace(again,"test","model",client=client)
        self.assertEqual(get(refreshed.sources,"N-DEMO-001").chunks["Customer CSV note"],"Changed interview note.")
        self.assertGreater(refreshed.ai_report["batches"],0)
        self.assertEqual(refreshed.ai_report["input_hash"],fingerprint(refreshed))
        original=again.model_dump()
        with self.assertRaisesRegex(ValueError,"citation/type"):
            analyze_workspace(again,"test","model",client=self.client(bad=True))
        self.assertEqual(again.model_dump(),original)

    def test_failure_and_unknown_synthesis_citation(self):
        from ai_workspace import analyze_workspace, BatchResult, Findings
        p=load_demo(); before=p.model_dump()
        client=Mock()
        client.responses.parse.return_value=SimpleNamespace(status="incomplete",output_parsed=None)
        with self.assertRaises(ValueError): analyze_workspace(p,"test","model",client=client)
        self.assertEqual(p.model_dump(),before)
        client=self.client(); original=client.responses.parse.side_effect
        def parse(**kw):
            result=original(**kw)
            if kw["text_format"] is Findings:
                result.output_parsed.findings[0].assessment_ids=["made-up"]
            return result
        client.responses.parse.side_effect=parse
        with self.assertRaisesRegex(ValueError,"unknown assessment"):
            analyze_workspace(p,"test","model",client=client)
        self.assertEqual(p.model_dump(),before)



class LocalProviderTests(unittest.TestCase):
    def test_local_links_attach_exact_passage_and_reject_unknown_ids(self):
        from local_ai import OllamaClient,LocalLinks
        from ai_workspace import BatchResult,parsed_call
        source={"id":"S1","kind":"Says","chunks":{"Paragraph 1":"The full original paragraph."}}
        payload={"sources":[source],"hypotheses":[{"id":"H3","statement":"Test"}]}
        result=LocalLinks(assessments=[dict(source_id="S1",hypothesis_id="H3",chunk="Paragraph 1",stance="Supports",rationale="Reason",limitations="Small sample")])
        with patch("local_ai.local_chat",return_value=result):
            parsed=parsed_call(OllamaClient(),"test",BatchResult,"Test",payload,1000)
            self.assertEqual(parsed.assessments[0].quote,"The full original paragraph.")
            result.assessments[0].hypothesis_id="invented"
            with self.assertRaisesRegex(ValueError,"unknown hypothesis"):
                parsed_call(OllamaClient(),"test",BatchResult,"Test",payload,1000)

    def test_cloud_model_is_not_local_fallback(self):
        from local_ai import local_chat,LocalLinks
        with self.assertRaisesRegex(ValueError,"not a cloud model"):
            local_chat("remote-cloud",LocalLinks,[])

    def test_demo_unmatched_and_reworded_hypothesis(self):
        from local_ai import DemoClient,scenarios
        from ai_workspace import analyze_workspace
        p=load_demo()
        p.customers[0].note="Unrelated fictional feedback that has no scenario match."
        updated=analyze_workspace(p,"demo","test",client=DemoClient(),provider="Demo scenarios")
        self.assertEqual(len(updated.assessments),len(p.assessments))
        p.customers[0].note=scenarios()[0]["text"]
        p.hypotheses[2].statement="A different hypothesis."
        updated=analyze_workspace(p,"demo","test",client=DemoClient(),provider="Demo scenarios")
        self.assertFalse(any(a.source_id=="N-DEMO-001" for a in updated.assessments))


class EvidenceTests(unittest.TestCase):
    def setUp(self):
        self.p=load_demo()

    def test_demo_and_roundtrip(self):
        self.assertEqual(len(self.p.customers),100)
        self.assertEqual(len(self.p.hypotheses),6)
        self.assertEqual(len([s for s in self.p.sources if s.kind=='Research']),4)
        restored=import_project(self.p.model_dump_json().encode())
        self.assertEqual(restored.model_dump(),self.p.model_dump())
        self.assertEqual(parse_customers(pd.read_csv(BytesIO(customer_frame(self.p.customers).to_csv(index=False).encode()))),self.p.customers)

    def test_source_edit_invalidates_approval_and_keeps_history(self):
        source=get(self.p.sources,'I-001')
        old=source.chunks.copy()
        edit_source(self.p,source.id,{'chunks':{'New passage':'A changed fictional interview.'},'expected_action':'','expected_by':None,'expectation_quote':'','expectation_reviewer':''})
        updated=get(self.p.sources,source.id)
        self.assertEqual(updated.history[0]['chunks'],old)
        self.assertEqual(updated.revision,2)
        self.assertFalse(any(a.source_id==source.id for a in current_assessments(self.p)))
        a=next(a for a in self.p.assessments if a.source_id==source.id)
        with self.assertRaises(ValueError):
            review_assessment(self.p,a.id,a.stance,a.quote,a.chunk,a.rationale,a.limitations,'Approved','Tester')
        review_assessment(self.p,a.id,'Unclear','A changed fictional interview.','New passage','Meaning changed.','Synthetic','Approved','Tester')
        self.assertTrue(any(a.source_id==source.id for a in current_assessments(self.p)))
        self.assertEqual(len(get(self.p.assessments,a.id).reviews),2)

    def test_pending_rejected_and_duplicate_links(self):
        s=get(self.p.sources,'I-012')
        self.assertFalse(any(a.source_id==s.id for a in current_assessments(self.p)))
        a=next(a for a in self.p.assessments if a.source_id==s.id)
        before=len(self.p.assessments)
        with self.assertRaises(ValueError):
            add_assessment(self.p,s.id,a.hypothesis_id,a.chunk,a.quote,a.stance,a.rationale)
        self.assertEqual(len(self.p.assessments),before)
        review_assessment(self.p,a.id,a.stance,a.quote,a.chunk,a.rationale,a.limitations,'Rejected','Tester')
        self.assertEqual(get(self.p.assessments,a.id).status,'Rejected')

    def test_counts_distinct_sources_not_excerpts(self):
        s=get(self.p.sources,'I-001')
        a=add_assessment(self.p,s.id,'H1','Interview paragraph 1','I do not have a regular doctor','Supports','Same interview.')
        review_assessment(self.p,a.id,a.stance,a.quote,a.chunk,a.rationale,'Duplicate context','Approved','Tester')
        self.assertEqual(triangulation(self.p).set_index('ID').loc['H1','Says: Supports'],1)

    def test_hypothesis_revision_preserves_original(self):
        original=get(self.p.hypotheses,'H1').statement
        revise_hypothesis(self.p,'H1','A revised testable statement.','Tester')
        self.assertEqual(get(self.p.hypotheses,'H1').history[0]['statement'],original)
        self.assertFalse(current_assessments(self.p,'H1'))

    def test_intention_checks_ignore_unknown_na_and_immature(self):
        self.assertEqual(len(contradictions(self.p)),2)
        c=get(self.p.customers,'DEMO-001')
        for status in ['Unknown','Not applicable']:
            c.events['consultation_booked']=Event(status=status)
            self.assertFalse(any(r['customer']==c.id for r in contradictions(self.p)))
        c.events['consultation_booked']=Event(status='Not completed')
        c.observed_through=date(2026,3,20)
        self.assertFalse(any(r['customer']==c.id for r in contradictions(self.p)))
        c.observed_through=date(2026,6,10)
        c.events['consultation_booked']=Event(status='Completed',on=date(2026,3,30))
        self.assertFalse(any(r['customer']==c.id for r in contradictions(self.p)))
        c.events['consultation_booked']=Event(status='Completed',on=date(2026,4,1))
        self.assertTrue(any(r['customer']==c.id for r in contradictions(self.p)))

    def test_opposing_links_same_customer_and_period(self):
        s=get(self.p.sources,'B-DEMO-001')
        # Align collection dates and create a different exact quote for a second assessment.
        s.source_date=date(2026,3,25)
        a=add_assessment(self.p,s.id,'H1','Customer action record','consultation_booked: Not completed','Contradicts','Test opposing-review mechanics.','Synthetic test only')
        review_assessment(self.p,a.id,a.stance,a.quote,a.chunk,a.rationale,a.limitations,'Approved','Tester')
        self.assertTrue(any(r['type']=='Opposing reviewed evidence' for r in contradictions(self.p)))
        s.context_version=2
        self.assertFalse(any(r['type']=='Opposing reviewed evidence' for r in contradictions(self.p)))

    def test_freshness_scoped_changes_and_dates(self):
        s=get(self.p.sources,'R-001')
        self.assertEqual(freshness(self.p,s,'H2',date(2026,9,21)),'Current')
        self.assertEqual(freshness(self.p,s,'H2',date(2027,1,1)),'Review due')
        self.p.changes.append(Change(id='CH-test',version=2,on=date(2026,9,1),description='Offer changed',hypotheses=['H2'],author='Tester'))
        self.p.context_version=2
        self.assertIn('Context changed',freshness(self.p,s,'H2',date(2026,9,21)))
        self.assertEqual(freshness(self.p,s,'H1',date(2026,9,21)),'Current')
        self.assertEqual(freshness(self.p,get(self.p.sources,'I-001'),'H1',date(2026,9,21)),'Review due')

    def test_import_rejects_broken_citations(self):
        raw=self.p.model_dump(mode='json')
        raw['assessments'][0]['quote']='This was never in the source.'
        with self.assertRaises(ValueError):
            import_project(json.dumps(raw).encode())
        with self.assertRaises(ValueError):
            import_project(b'{bad json')

    def test_customer_import_atomic_and_updates_source(self):
        c=self.p.customers[0].model_copy(deep=True)
        with self.assertRaises(ValueError):
            save_customers(self.p,[c])
        c.number=2
        before=self.p.model_dump_json()
        with self.assertRaises(ValueError):
            save_customers(self.p,[c],True)
        self.assertEqual(before,self.p.model_dump_json())
        c.number=1
        c.satisfaction=1
        save_customers(self.p,[c],True)
        self.assertEqual(get(self.p.sources,'B-DEMO-001').revision,2)
        self.assertIn('Satisfaction: 1/5',get(self.p.sources,'B-DEMO-001').chunks['Customer action record'])

    def test_ai_contract_and_exact_citation_validation(self):
        s=get(self.p.sources,'I-012')
        proposal=Proposals(assessments=[dict(hypothesis_id='H2',chunk='Interview paragraph 1',quote='Separate appointments fit my flexible schedule.',stance='Unclear',rationale='Different cause mentioned.',limitations='One fictional customer.')],note='Needs human review.')
        mock=Mock()
        mock.responses.parse.return_value=SimpleNamespace(output_parsed=proposal,status='completed')
        result=ai_proposals(self.p,s.id,['Interview paragraph 1'],'fake','test-model',mock)
        kwargs=mock.responses.parse.call_args.kwargs
        self.assertFalse(kwargs['store'])
        self.assertNotIn('customers',json.loads(kwargs['input'][1]['content']))
        self.assertEqual(accept_proposals(self.p,s.id,result,'Mock AI test'),1)
        self.assertEqual(self.p.assessments[-1].status,'Pending')
        count=len(self.p.assessments)
        proposal.assessments[0].quote='invented quote'
        with self.assertRaises(ValueError):
            accept_proposals(self.p,s.id,proposal,'Mock AI test')
        self.assertEqual(count,len(self.p.assessments))
        mock.responses.parse.return_value=SimpleNamespace(output_parsed=None,status='completed')
        with self.assertRaises(ValueError):
            ai_proposals(self.p,s.id,['Interview paragraph 1'],'fake','test-model',mock)

    def test_document_extraction(self):
        self.assertEqual(extract_document('note.txt',b'First paragraph.\n\nSecond paragraph.')['Paragraph 2'],'Second paragraph.')
        from docx import Document
        doc=Document()
        doc.add_paragraph('Fictional customer interview.')
        doc.add_table(rows=1,cols=1).cell(0,0).text='Evidence in a table'
        buf=BytesIO(); doc.save(buf)
        chunks=extract_document('note.docx',buf.getvalue())
        self.assertEqual(chunks['Table 1, row 1'],'Evidence in a table')
        from pypdf import PdfWriter
        buf=BytesIO(); writer=PdfWriter(); writer.add_blank_page(width=100,height=100); writer.write(buf)
        with self.assertRaisesRegex(ValueError,'without extractable text'):
            extract_document('scan.pdf',buf.getvalue())

    def test_age_scope_preserves_research_but_excludes_other_customers(self):
        scoped=current_assessments(self.p,customer_ids=set())
        self.assertEqual(len(scoped),4)
        self.assertTrue(all(get(self.p.sources,a.source_id).kind=='Research' for a in scoped))

    def test_proposal_batch_invalid_citation_is_atomic(self):
        proposals=Proposals(assessments=[
            dict(hypothesis_id='H2',chunk='Interview paragraph 1',quote='Separate appointments fit my flexible schedule.',stance='Unclear',rationale='Test',limitations='Test'),
            dict(hypothesis_id='H3',chunk='Interview paragraph 1',quote='Fabricated.',stance='Supports',rationale='Test',limitations='Test')],note='')
        before=len(self.p.assessments)
        with self.assertRaises(ValueError):
            accept_proposals(self.p,'I-012',proposals,'Test')
        self.assertEqual(before,len(self.p.assessments))

    def test_demo_has_distinct_current_stories(self):
        for hid,action in [("H3","Double down"),("H4","Revise"),("H5","Retire")]:
            result=hypothesis_thesis(self.p,hid)
            self.assertEqual(result["action"],action)
            self.assertFalse(result["stale_sources"])
        self.assertEqual(hypothesis_thesis(self.p,"H1")["recommendation"],"Revalidate first")
        interview=get(self.p.sources,"I-013")
        self.assertEqual(len(interview.chunks),3)
        self.assertEqual({a.hypothesis_id for a in current_assessments(self.p) if a.source_id==interview.id},{"H3","H4"})
        self.assertFalse(self.p.decisions)
        self.assertFalse(any("comparison" in c.segment for c in self.p.customers))

    def test_activity_windows_deduplicate_and_measure_churn(self):
        self.p.activity_start=date(2026,3,1)
        self.p.activity_end=date(2026,4,29)
        self.p.activity_events=[ActivityEvent(customer_id=cid,on=date.fromisoformat(day),activity=action) for cid,day,action in [
            ("DEMO-001","2026-03-05","App opened"),("DEMO-001","2026-03-05","Results opened"),
            ("DEMO-002","2026-03-06","App opened"),("DEMO-001","2026-04-15","App opened"),
            ("DEMO-003","2026-04-20","Care plan viewed")]]
        report=activity_report(self.p,date(2026,4,29))
        self.assertEqual(report["mau"],2)
        self.assertEqual(report["previous_mau"],2)
        self.assertEqual(report["churn"],.5)
        self.assertEqual(report["retention"],.5)
        self.assertEqual(report["growth"],0)
        self.assertEqual(report["lost"],{"DEMO-002"})
        self.assertEqual(report["newly_active"],{"DEMO-003"})
        daily=report["daily"].set_index("Date")
        self.assertEqual(daily.loc[pd.Timestamp("2026-03-05"),"Daily active users"],1)
        self.assertEqual(report["dau"],0)
        self.assertIsNone(activity_report(self.p,date(2026,3,20))["mau"])
        self.assertIsNone(activity_report(self.p,date(2026,4,29),set())["churn"])
        self.assertEqual(report["monthly"].iloc[-1].Coverage,"Partial")
        with self.assertRaises(ValueError):
            activity_report(self.p,date(2026,4,30))
        self.assertEqual(len(import_project(self.p.model_dump_json().encode()).activity_events),5)

    def test_activity_validation_and_old_project_compatibility(self):
        raw=self.p.model_dump(mode="json")
        for name in ["activity_events","activity_start","activity_end"]:
            raw.pop(name)
        old=Project.model_validate(raw)
        self.assertFalse(old.activity_events)
        raw.update(activity_start="2026-03-01",activity_end="2026-04-01",activity_events=[
            dict(customer_id="unknown",on="2026-03-05",activity="App opened")])
        with self.assertRaises(ValueError):
            Project.model_validate(raw)

    def test_ai_passage_type_remains_pending_until_review(self):
        s=get(self.p.sources,"I-012")
        proposal=Proposals(assessments=[dict(hypothesis_id="H4",chunk="Interview paragraph 1",
            quote=s.chunks["Interview paragraph 1"],stance="Unclear",rationale="Different barrier.",limitations="One account.",evidence_kind="Says")],note="")
        accept_proposals(self.p,s.id,proposal,"Mock AI")
        a=self.p.assessments[-1]
        self.assertEqual(a.evidence_kind,"Says")
        self.assertEqual(a.status,"Pending")
        source=get(self.p.sources,"R-001")
        a=add_assessment(self.p,source.id,"H1",next(iter(source.chunks)),next(iter(source.chunks.values())),"Unclear","Background",evidence_kind="Does")
        with self.assertRaises(ValueError):
            review_assessment(self.p,a.id,a.stance,a.quote,a.chunk,a.rationale,"","Approved","Tester")
        review_assessment(self.p,a.id,a.stance,a.quote,a.chunk,a.rationale,"","Approved","Tester","Research")
        self.assertEqual(get(self.p.assessments,a.id).evidence_kind,"Research")

    def thesis_fixture(self, stance):
        self.p.sources=[]
        self.p.assessments=[]
        get(self.p.hypotheses,'H1').review_on=date.today()+timedelta(days=30)
        for n in [1,2,11,12,13]:
            for kind in ['Says','Does']:
                s=Source(id=f'{kind}-{n}',title='Test source',kind=kind,customer_id=f'DEMO-{n:03d}',
                    chunks={'p':'Exact test evidence.'},source_date=date.today(),review_on=date.today()+timedelta(days=30),context='Unit test')
                self.p.sources.append(s)
                a=add_assessment(self.p,s.id,'H1','p','Exact test evidence.',stance,'Test rationale')
                review_assessment(self.p,a.id,a.stance,a.quote,a.chunk,a.rationale,'Test','Approved','Tester')

    def test_thesis_requires_corroboration_and_repeated_cohorts(self):
        self.thesis_fixture('Supports')
        result=hypothesis_thesis(self.p,'H1')
        self.assertEqual(result['action'],'Double down')
        self.assertEqual(result['supporting_customers'],5)
        self.assertEqual(result['supporting_cohorts'],2)
        self.assertEqual(len(result['assessment_ids']),10)
        self.assertFalse(self.p.decisions)
        for s in self.p.sources:
            if s.kind=='Does': s.archived=True
        self.assertEqual(hypothesis_thesis(self.p,'H1')['action'],'Keep testing')

    def test_retirement_requires_current_customer_counterevidence(self):
        self.thesis_fixture('Contradicts')
        self.assertEqual(hypothesis_thesis(self.p,'H1')['recommendation'],'Consider retiring')
        self.p.sources[0].review_on=date.today()
        self.assertEqual(hypothesis_thesis(self.p,'H1')['recommendation'],'Revalidate first')
        self.assertFalse(self.p.decisions)

    def test_research_and_no_evidence_cannot_trigger_retirement(self):
        self.thesis_fixture('Contradicts')
        for s in self.p.sources:
            s.kind='Research'; s.customer_id=''
        for a in self.p.assessments: a.evidence_kind='Research'
        self.assertEqual(hypothesis_thesis(self.p,'H1')['action'],'Keep testing')
        self.p.assessments=[]
        self.assertEqual(hypothesis_thesis(self.p,'H1')['status'],'Limited evidence')

    def test_thesis_scope_and_duplicate_customer_safeguards(self):
        self.thesis_fixture('Supports')
        self.assertEqual(hypothesis_thesis(self.p,'H1',{'DEMO-001'})['action'],'Keep testing')
        for s in self.p.sources: s.customer_id='DEMO-001'
        result=hypothesis_thesis(self.p,'H1')
        self.assertEqual(result['supporting_customers'],1)
        self.assertEqual(result['action'],'Keep testing')


def widget(items,label):
    return next(x for x in items if x.label==label)


class AppTests(unittest.TestCase):
    def setUp(self):
        self.at=AppTest.from_file('../app.py').run(timeout=30)
        widget(self.at.radio,'Workspace').set_value('Hypotheses').run(timeout=30)

    def page(self,label):
        widget(self.at.radio,'Workspace').set_value(label).run(timeout=30)
        self.assertFalse(self.at.exception)

    def test_sidebar_and_numeric_cohort_order(self):
        self.assertEqual(widget(self.at.sidebar.radio,"Workspace").value,"Hypotheses")
        self.page("Customers")
        self.assertEqual(widget(self.at.multiselect,"Cohort").options,["1–10","11–30","31–100"])

    def test_filter_empty_states_and_multiple_passages(self):
        self.page("Evidence")
        widget(self.at.selectbox,"Sort evidence").set_value("Most passages").run()
        self.assertFalse(self.at.exception)
        sid=widget(self.at.selectbox,"Choose evidence").value
        self.assertGreater(len(get(self.at.session_state.project.sources,sid).chunks),1)
        self.assertGreater(len(widget(self.at.selectbox,"Source passage").options),1)
        widget(self.at.text_input,"Find evidence").set_value("no-such-evidence").run()
        self.assertTrue(any("No matches" in x.value for x in self.at.info))
        self.page("Hypotheses")
        widget(self.at.multiselect,"Recommendations").set_value(["Consider retiring"]).run()
        self.assertEqual(widget(self.at.selectbox,"Choose hypothesis").value,"H5")
        widget(self.at.text_input,"Find hypothesis").set_value("no-such-hypothesis").run()
        self.assertTrue(any("No matching" in x.value for x in self.at.info))
        self.page("Customers")
        widget(self.at.text_input,"Find customer").set_value("DEMO-013").run()
        self.assertEqual(widget(self.at.selectbox,"Choose customer").value,"DEMO-013")
        widget(self.at.text_input,"Find customer").set_value("no-such-customer").run()
        self.assertTrue(any("0 customers in view" in x.value for x in self.at.caption))
        self.assertFalse(self.at.exception)

    def test_addition_sort_is_distinct_from_source_dates(self):
        self.page("Evidence")
        self.assertFalse(any(x.label=="Evidence filters" for x in self.at.expander))
        newest=self.at.session_state.project.sources[-1].id
        self.assertEqual(widget(self.at.selectbox,"Choose evidence").options[0].split(" · ")[0],newest)
        widget(self.at.selectbox,"Sort evidence").set_value("Oldest added").run()
        self.assertEqual(widget(self.at.selectbox,"Choose evidence").options[0].split(" · ")[0],self.at.session_state.project.sources[0].id)
        self.page("Customers")
        self.assertEqual(widget(self.at.selectbox,"Choose customer").options[0],"DEMO-100")
        widget(self.at.selectbox,"Sort customers").set_value("Oldest added").run()
        self.assertEqual(widget(self.at.selectbox,"Choose customer").options[0],"DEMO-001")
        self.assertTrue({"Segment","Channel","Cohort"}<={x.label for x in self.at.multiselect})
        self.assertFalse(self.at.exception)

    def test_dashboard_filters_and_non_destructive_activity_load(self):
        self.page("Dashboard")
        self.assertEqual(len(self.at.metric),6)
        widget(self.at.multiselect,"Segments").set_value(["Young professional / wellness"]).run()
        self.assertFalse(self.at.exception)
        p=self.at.session_state.project
        p.activity_events=[];p.activity_start=None;p.activity_end=None
        before=p.assessments[0].model_dump()
        self.at.run()
        widget(self.at.button,"Load demo activity").click().run()
        self.assertTrue(self.at.session_state.project.activity_events)
        self.assertEqual(self.at.session_state.project.assessments[0].model_dump(),before)
        self.assertFalse(self.at.exception)

    def test_workspace_ai_controls_refresh_and_review(self):
        self.at.session_state.analysis_mode="OpenAI"
        from ai_workspace import analyze_workspace
        self.at.session_state.ai_session_key="test"
        self.at.session_state.ai_session_model="model"
        self.page("AI analysis")
        self.assertTrue(widget(self.at.button,"Analyze workspace").disabled)
        widget(self.at.checkbox,"Send workspace data to OpenAI for analysis").check().run()
        widget(self.at.checkbox,"Automatically update after saved data changes").check().run()
        client=WorkspaceAITests().client()
        with patch("openai.OpenAI",return_value=client):
            widget(self.at.button,"Analyze workspace").click().run(timeout=60)
            self.assertFalse(self.at.exception)
            self.assertTrue(self.at.session_state.project.ai_report)
            calls=client.responses.parse.call_count
            self.page("Customers")
            self.assertEqual(client.responses.parse.call_count,calls)
            self.at.session_state.project.customers[0].note="A newly saved input."
            self.page("AI analysis")
            self.assertGreater(client.responses.parse.call_count,calls)
            self.assertFalse(self.at.exception)
            current=self.at.session_state.project.model_dump()
            self.at.session_state.project.customers[0].note="Another input, provider unavailable."
            client.responses.parse.side_effect=RuntimeError("Provider unavailable")
            self.at.run()
            failed_calls=client.responses.parse.call_count
            self.assertTrue(self.at.session_state.ai_batch_error)
            self.at.run()
            self.assertEqual(client.responses.parse.call_count,failed_calls)
            self.assertEqual(self.at.session_state.project.ai_report,current["ai_report"])
            widget(self.at.button,"Review selected proposal").click().run()
            self.assertEqual(widget(self.at.radio,"Workspace").value,"Evidence")
            self.assertFalse(self.at.exception)

    def test_demo_scenarios_without_network(self):
        from local_ai import scenarios
        self.page("AI analysis")
        self.assertEqual(widget(self.at.radio,"Analysis mode").value,"Demo scenarios")
        self.assertFalse(widget(self.at.button,"Analyze workspace").disabled)
        with patch("urllib.request.urlopen",side_effect=AssertionError("Demo must not call network")), patch("openai.OpenAI",side_effect=AssertionError("Demo must not call API")):
            widget(self.at.button,"Analyze workspace").click().run()
            self.assertFalse(self.at.exception)
            baseline=len(self.at.session_state.project.assessments)
            self.at.session_state.project.customers[0].note=scenarios()[1]["text"]
            widget(self.at.button,"Analyze workspace").click().run()
            self.assertFalse(self.at.exception)
            project=self.at.session_state.project
            a=[a for a in project.assessments if a.source_id=="N-DEMO-001"][-1]
            self.assertEqual(a.hypothesis_id,"H4")
            self.assertEqual(a.stance,"Contradicts")
            self.assertEqual(a.status,"Pending")
            self.assertEqual(project.ai_report["provider"],"Demo scenarios")
            self.assertGreater(len(project.assessments),baseline)
            self.assertFalse(project.decisions)
            self.assertTrue(any("affordability" in f["conclusion"] for f in project.ai_report["findings"]))

    def test_all_screens(self):
        for page in ['Dashboard','Evidence','Customers','Contradictions','Decisions','AI analysis','AI setup','Hypotheses']:
            self.page(page)

    def test_add_edit_evidence_and_review(self):
        self.page('Evidence')
        # The add-source editor precedes the selected-source editor.
        [x for x in self.at.text_input if x.label=='Evidence title'][0].set_value('Test fictional interview')
        [x for x in self.at.text_area if x.label=='Source text'][0].set_value('I have no regular doctor.')
        widget(self.at.button,'Add evidence').click().run()
        self.assertFalse(self.at.exception)
        p=self.at.session_state.project
        s=next(x for x in p.sources if x.title=='Test fictional interview')
        widget(self.at.selectbox,'Choose evidence').set_value(s.id).run()
        widget(self.at.text_area,'Why this position?').set_value('Directly reports lack of regular care.')
        widget(self.at.button,'Save proposed link').click().run()
        widget(self.at.text_input,'Reviewer name').set_value('Test reviewer')
        widget(self.at.button,'Save classification review').click().run()
        a=next(x for x in self.at.session_state.project.assessments if x.source_id==s.id)
        self.assertEqual(a.status,'Approved')
        [x for x in self.at.text_area if x.label=='Source text'][-1].set_value('I now have a regular doctor.')
        widget(self.at.button,'Save evidence changes').click().run()
        self.assertEqual(get(self.at.session_state.project.assessments,a.id).status,'Needs review')
        self.assertFalse(self.at.exception)

    def test_customer_add_and_project_import_failure(self):
        self.page('Customers')
        widget(self.at.button,'Save customer').click().run()
        self.assertEqual(len(self.at.session_state.project.customers),101)
        self.assertFalse(self.at.exception)
        widget(self.at.file_uploader,"Restore exported project").upload('bad.json',b'{oops','application/json').run()
        widget(self.at.button,'Import project').click().run()
        self.assertEqual(len(self.at.session_state.project.customers),101)
        self.assertTrue(self.at.error)

    def test_change_and_decision(self):
        self.page('Decisions')
        widget(self.at.text_area,'What changed, and which earlier findings might no longer apply?').set_value('Booking now offers evening appointments.')
        widget(self.at.multiselect,'Hypotheses to revalidate').set_value(['H6'])
        widget(self.at.text_input,'Change author').set_value('Tester')
        widget(self.at.button,'Log change and flag evidence').click().run()
        self.assertEqual(self.at.session_state.project.context_version,2)
        widget(self.at.text_area,'Decision rationale').set_value('Need interviews under the new offer.')
        widget(self.at.text_area,'Next test / evidence to collect').set_value('Interview five customers.')
        widget(self.at.text_input,'Decision author').set_value('Tester')
        widget(self.at.button,'Save human decision').click().run()
        self.assertEqual(len(self.at.session_state.project.decisions),1)
        self.assertFalse(self.at.exception)

    def test_project_restore_and_document_upload(self):
        exported=self.at.session_state.project.model_dump_json().encode()
        widget(self.at.file_uploader,"Restore exported project").upload('project.json',exported,'application/json').run()
        widget(self.at.button,'Import project').click().run()
        self.assertFalse(self.at.exception)
        self.assertEqual(len(self.at.session_state.project.customers),100)
        self.page('Evidence')
        widget(self.at.file_uploader,'Interview or research document').upload('sample.txt',b'I need a doctor.\n\nCost is my barrier.','text/plain').run()
        widget(self.at.button,'Extract document for review').click().run()
        widget(self.at.button,'Add evidence').click().run()
        self.assertFalse(self.at.exception)
        source=next(s for s in self.at.session_state.project.sources if s.title=='sample.txt')
        self.assertEqual(source.chunks['Paragraph 2'],'Cost is my barrier.')

    def test_freshness_review_and_contradiction_resolution(self):
        self.page('Evidence')
        widget(self.at.selectbox,'Choose evidence').set_value('I-001').run()
        widget(self.at.text_area,'Freshness review rationale').set_value('Rechecked the context; still relevant.')
        widget(self.at.text_input,'Freshness reviewer').set_value('Tester')
        widget(self.at.button,'Save freshness review').click().run()
        self.assertFalse(self.at.exception)
        s=get(self.at.session_state.project.sources,'I-001')
        self.assertEqual(freshness(self.at.session_state.project,s),'Current')
        self.assertEqual(len(s.context_reviews),1)
        self.page('Contradictions')
        widget(self.at.text_area,'Explanation or next question').set_value('Ask whether booking hours were available.')
        widget(self.at.text_input,'Mismatch reviewer').set_value('Tester')
        widget(self.at.button,'Save mismatch review').click().run()
        self.assertTrue(self.at.session_state.project.contradiction_reviews)
        self.assertFalse(self.at.exception)

    def test_csv_update_duplicates_and_threshold(self):
        self.page('Customers')
        sample=customer_frame(self.at.session_state.project.customers[:1])
        sample['satisfaction']=1
        widget(self.at.file_uploader,'Customer CSV').upload('customer.csv',sample.to_csv(index=False).encode(),'text/csv').run()
        widget(self.at.button,'Apply customer CSV').click().run()
        self.assertTrue(self.at.error)
        widget(self.at.radio,'Matching IDs').set_value('Update matching IDs').run()
        widget(self.at.button,'Apply customer CSV').click().run()
        self.assertEqual(self.at.session_state.project.customers[0].satisfaction,1)
        self.assertEqual(len(self.at.session_state.project.customers),100)
        widget(self.at.number_input,'Power-user score cutoff').set_value(11).run()
        self.assertTrue(any('0 meet' in x.value for x in self.at.caption))
        self.assertFalse(self.at.exception)

    def test_review_history_keeps_selected_link_and_latest_review(self):
        self.page("Evidence")
        widget(self.at.selectbox,"Choose evidence").set_value("I-013").run()
        links=[a for a in self.at.session_state.project.assessments if a.source_id=="I-013"]
        aid=links[1].id
        widget(self.at.selectbox,"Choose link").set_value(aid).run()
        widget(self.at.text_input,"Reviewer name").set_value("History reviewer")
        widget(self.at.radio,"Review outcome").set_value("Rejected")
        widget(self.at.button,"Save classification review").click().run()
        self.assertFalse(self.at.exception)
        self.assertEqual(widget(self.at.selectbox,"Choose link").value,aid)
        a=get(self.at.session_state.project.assessments,aid)
        self.assertEqual(a.reviews[-1]["reviewer"],"History reviewer")
        histories=[x.value for x in self.at.dataframe if "Reviewed at (UTC)" in x.value.columns]
        self.assertTrue(any(((h["Reviewer"]=="History reviewer") & (h["Outcome"]=="Rejected")).any() for h in histories))
        widget(self.at.text_input,"Reviewer name").set_value("Second reviewer")
        widget(self.at.radio,"Review outcome").set_value("Approved")
        widget(self.at.button,"Save classification review").click().run()
        self.assertEqual(widget(self.at.selectbox,"Choose link").value,aid)
        self.assertEqual(len(get(self.at.session_state.project.assessments,aid).reviews),3)
        restored=import_project(self.at.session_state.project.model_dump_json().encode())
        self.assertEqual(get(restored.assessments,aid).reviews[-1]["reviewer"],"Second reviewer")

    def test_approved_review_visible_in_recent_history(self):
        self.page("Evidence")
        widget(self.at.selectbox,"Choose evidence").set_value("I-011").run()
        widget(self.at.text_input,"Reviewer name").set_value("Filtered reviewer")
        widget(self.at.button,"Save classification review").click().run()
        self.assertFalse(self.at.exception)
        histories=[x.value for x in self.at.dataframe if "Reviewed at (UTC)" in x.value.columns]
        self.assertTrue(any((h["Reviewer"]=="Filtered reviewer").any() for h in histories))

    def test_ai_setup_is_session_only_and_connection_check_is_explicit(self):
        self.at.session_state.analysis_mode="OpenAI"
        self.page("AI setup")
        widget(self.at.text_input,"OpenAI API key").set_value("test-key-not-real")
        widget(self.at.text_input,"Model ID").set_value("test-model")
        with patch("openai.OpenAI") as provider:
            widget(self.at.button,"Use key for this session").click().run()
            provider.assert_not_called()
            self.assertEqual(self.at.session_state.ai_session_key,"test-key-not-real")
            self.assertEqual(widget(self.at.text_input,"OpenAI API key").value,"")
            self.assertNotIn("test-key-not-real",self.at.session_state.project.model_dump_json())
            widget(self.at.button,"Test API connection").click().run()
            provider.return_value.__enter__.return_value.models.retrieve.assert_called_once_with("test-model")
            self.assertTrue(self.at.session_state.ai_connection_result[0])
        with patch("openai.OpenAI",side_effect=RuntimeError("test-key-not-real should never be echoed")):
            widget(self.at.button,"Test API connection").click().run()
            self.assertFalse(self.at.session_state.ai_connection_result[0])
            self.assertFalse(any("test-key-not-real" in x.value for x in self.at.error))
        widget(self.at.button,"Go to Evidence").click().run()
        self.assertEqual(widget(self.at.radio,"Workspace").value,"Evidence")
        widget(self.at.checkbox,"Send these selected passages for AI analysis").check().run()
        self.assertFalse(widget(self.at.button,"Analyze selected evidence").disabled)
        self.page("AI setup")
        widget(self.at.button,"Remove session key").click().run()
        self.assertNotIn("ai_session_key",self.at.session_state)
        self.assertFalse(self.at.exception)

    def test_approved_evidence_handoff_preselects_citation(self):
        self.page("Evidence")
        widget(self.at.selectbox,"Choose evidence").set_value("I-013").run()
        aid=widget(self.at.selectbox,"Choose link").value
        widget(self.at.button,"Use this evidence in a decision").click().run()
        self.assertEqual(widget(self.at.radio,"Workspace").value,"Decisions")
        self.assertEqual(widget(self.at.multiselect,"Cite reviewed assessments").value,[aid])
        self.assertFalse(self.at.session_state.project.decisions)
        widget(self.at.text_area,"Decision rationale").set_value("Review the paired observation before broadening recruitment.")
        widget(self.at.text_area,"Next test / evidence to collect").set_value("Recruit five more wellness customers.")
        widget(self.at.text_input,"Decision author").set_value("Test reviewer")
        widget(self.at.button,"Save human decision").click().run()
        self.assertEqual(self.at.session_state.project.decisions[-1].assessment_ids,[aid])
        self.assertFalse(self.at.exception)

    def test_source_category_optional_notes_and_automatic_quote(self):
        self.page("Evidence")
        widget(self.at.selectbox,"Source category").set_value("Follow-up interview").run()
        [x for x in self.at.text_input if x.label=="Evidence title"][0].set_value("Follow-up without extra notes")
        [x for x in self.at.text_area if x.label=="Source text"][0].set_value("I need help interpreting my report.")
        widget(self.at.button,"Add evidence").click().run()
        self.assertFalse(self.at.exception)
        sid=widget(self.at.selectbox,"Choose evidence").value
        source=get(self.at.session_state.project.sources,sid)
        self.assertEqual(source.category,"Follow-up interview")
        self.assertEqual(source.kind,"Says")
        self.assertEqual(source.context,"")
        self.assertFalse(any(x.label=="Exact quotation" for x in self.at.text_area))
        widget(self.at.text_area,"Why this position?").set_value("The customer reports an interpretation need.")
        widget(self.at.button,"Save proposed link").click().run()
        a=next(a for a in self.at.session_state.project.assessments if a.source_id==sid)
        self.assertEqual(a.quote,source.chunks[a.chunk])
        self.assertEqual(a.status,"Pending")
        widget(self.at.text_input,"Reviewer name").set_value("Tester")
        widget(self.at.button,"Save classification review").click().run()
        self.assertEqual(get(self.at.session_state.project.assessments,a.id).status,"Approved")
        self.assertFalse(self.at.exception)

    def test_thesis_to_human_decision(self):
        widget(self.at.button,'Review recommendation').click().run()
        self.assertFalse(self.at.exception)
        self.assertEqual(widget(self.at.radio,'Workspace').value,'Decisions')
        self.assertFalse(self.at.session_state.project.decisions)
        self.assertTrue(widget(self.at.text_area,'Decision rationale').value)
        widget(self.at.text_input,'Decision author').set_value('Reviewer')
        widget(self.at.button,'Save human decision').click().run()
        self.assertFalse(self.at.exception)
        decision=self.at.session_state.project.decisions[-1]
        self.assertEqual(decision.author,'Reviewer')
        self.assertEqual(decision.recommendation_snapshot['method'],'Review rules v1')


if __name__=='__main__':
    unittest.main()


class ReadableSummaryTests(unittest.TestCase):
    def test_grounded_distinct_findings_and_data_changes(self):
        from ai_workspace import analyze_workspace
        from local_ai import DemoClient
        p=analyze_workspace(load_demo(),"demo","Prepared scenarios v1",client=DemoClient(),provider="Demo scenarios")
        findings=p.ai_report["findings"]
        self.assertGreater(len({f["conclusion"] for f in findings}),10)
        self.assertFalse(any("Prepared baseline summary" in f["conclusion"] for f in findings))
        h3=next(f for f in findings if f["target"]=="H3")
        self.assertIn("Consider doubling down",h3["conclusion"])
        self.assertIn("supporting",h3["conclusion"])
        for source in p.sources:
            source.review_on=date.today()-timedelta(days=1)
        changed=analyze_workspace(p,"demo","Prepared scenarios v1",client=DemoClient(),provider="Demo scenarios")
        self.assertIn("Revalidate first",next(f for f in changed.ai_report["findings"] if f["target"]=="H3")["conclusion"])
        self.assertEqual(p.decisions,changed.decisions)

    def test_visible_summaries_and_old_report_upgrade(self):
        from ai_workspace import analyze_workspace
        from local_ai import DemoClient
        p=analyze_workspace(load_demo(),"demo","Prepared scenarios v1",client=DemoClient(),provider="Demo scenarios")
        p.ai_report.pop("summary_version")
        before=p.model_dump(exclude={"ai_report"})
        app=AppTest.from_file("../app.py",default_timeout=30)
        app.session_state.project=p
        app.run()
        self.assertFalse(app.exception)
        self.assertEqual(app.session_state.project.model_dump(exclude={"ai_report"}),before)
        for page in ["Dashboard","Hypotheses","Customers","Evidence","Contradictions","Decisions","AI analysis"]:
            app.sidebar.radio[0].set_value(page).run()
            self.assertFalse(app.exception)
            self.assertIn("Preliminary findings" if page=="AI analysis" else "AI summary",[h.value for h in app.subheader])


class DataWorkbenchTests(unittest.TestCase):
    def test_evidence_scope_and_empty_selection(self):
        from data_workbench import query_data
        p=load_demo(); before=p.model_dump()
        args=[p,"Evidence comparison",["Does","Research"],[],["H3"],["Approved"],True,date(2020,1,1),date(2030,1,1)]
        rows,counts,facts=query_data(*args)
        self.assertFalse(rows.empty)
        self.assertTrue(set(rows.Type)<={"Does","Research"})
        self.assertEqual(set(rows.Hypothesis),{"H3"})
        self.assertEqual(set(rows.Review),{"Approved"})
        self.assertEqual(set(rows.Freshness),{"Current"})
        self.assertEqual(counts.Links.sum(),len(rows))
        args[2]=[]
        self.assertTrue(query_data(*args)[0].empty)
        self.assertEqual(before,p.model_dump())

    def test_activity_unique_customers_and_range(self):
        from data_workbench import query_data
        p=load_demo(); day=p.activity_end; segment=p.customers[0].segment
        frame,_,_=query_data(p,"Daily active users",[],[segment],[],[],True,day,day)
        ids={c.id for c in p.customers if c.segment==segment}
        expected=len({e.customer_id for e in p.activity_events if e.on==day and e.customer_id in ids})
        self.assertEqual(len(frame),1)
        self.assertEqual(frame.iloc[0]["Daily active users"],expected)
        with self.assertRaises(ValueError):query_data(p,"Daily active users",[],[],[],[],True,day,day+timedelta(days=1))

    def test_citation_validation_and_local_scope(self):
        from data_workbench import interpret,Answer
        frame=pd.DataFrame([dict(Reference="A1",Type="Does",Quote="Observed action")])
        with patch("local_ai.local_chat",return_value=Answer(answer="One observed action.",references=["A1"],limitations="One record.")) as model:
            result=interpret("What happened?",frame,"1 action","Local Ollama","test")
            self.assertEqual(result["references"],["A1"])
            payload=json.loads(model.call_args.args[2][1]["content"])
            self.assertEqual(payload["rows"][0]["Type"],"Does")
        with patch("local_ai.local_chat",return_value=Answer(answer="Bad citation",references=["FAKE"],limitations="test")):
            with self.assertRaises(ValueError):interpret("Question",frame,"1 action","Local Ollama","test")

    def test_workbench_ui_and_stale_result(self):
        app=AppTest.from_file("../app.py",default_timeout=30).run()
        app.sidebar.radio[0].set_value("Ask your data").run()
        self.assertFalse(app.exception)
        next(b for b in app.button if b.label=="Run analysis").click().run()
        self.assertFalse(app.exception)
        self.assertIn("Result",[h.value for h in app.subheader])
        self.assertNotIn("Says",set(app.session_state.ask_result["frame"].Type))
        next(s for s in app.selectbox if s.label=="Analyze").set_value("Daily active users").run()
        self.assertTrue(any("scope changed" in i.value for i in app.info))
        next(b for b in app.button if b.label=="Run analysis").click().run()
        self.assertFalse(app.exception)
        self.assertIn("Daily active users",app.session_state.ask_result["frame"].columns)
